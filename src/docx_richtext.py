# -*- coding: utf-8 -*-
"""DOCX 富文本处理：下标标记、上标交叉引用（GB/T 7714 顺序编码制）。

- 正文中的 ~x~ 渲染为 Word 原生下标（如 n~H~ -> nH）
- 正文中的 [n] / [n,m] 渲染为上标交叉引用域，与文末条目联动
"""
from __future__ import annotations

import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# 单个 [数字] 或 [数字,数字]
CITE_SPLIT = re.compile(r"(\[\d+(?:,\d+)*\])")
CITE_FULL = re.compile(r"^\[([\d,]+)\]$")
REF_HEAD = re.compile(r"^\[(\d+)\]\s*")

_id_counter = [1000]


def add_bookmark(par, name):
    """返回 (bookmarkStart, bookmarkEnd) 两个 XML 元素。"""
    _id_counter[0] += 1
    bid = str(_id_counter[0])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    return start, end


def add_ref_field(par, bookmark, cached, superscript=True):
    """插入 REF 交叉引用域；域带缓存结果，未更新域时也能正常显示。"""
    def _run():
        r = par.add_run()
        r.font.superscript = superscript
        return r

    r = _run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r._r.append(fc)

    r = _run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " REF " + bookmark + " " + chr(92) + "h "
    r._r.append(instr)

    r = _run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r._r.append(fc)

    cr = par.add_run(cached)
    cr.font.superscript = superscript

    r = _run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r._r.append(fc)


def add_rich_runs(par, text, size):
    """把 ~下标~ 与 [引用] 解析为对应的 Word 格式。"""
    for i, part in enumerate(re.split(r"~([^~]+)~", text)):
        if not part:
            continue
        if i % 2 == 1:                      # ~x~ -> 下标
            run = par.add_run(part)
            run.font.size = size
            run.font.subscript = True
            continue
        for seg in CITE_SPLIT.split(part):
            if not seg:
                continue
            m = CITE_FULL.match(seg)
            if not m:                        # 普通文字
                par.add_run(seg).font.size = size
                continue
            nums = [x.strip() for x in m.group(1).split(",")]
            for pos, n in enumerate(nums):
                if pos == 0:
                    b = par.add_run("[")
                    b.font.size = size
                    b.font.superscript = True
                else:
                    sep = par.add_run(",")
                    sep.font.size = size
                    sep.font.superscript = True
                add_ref_field(par, "_Ref" + n, n)
            b = par.add_run("]")
            b.font.size = size
            b.font.superscript = True
    return par
