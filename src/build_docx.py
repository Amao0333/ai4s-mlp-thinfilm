# -*- coding: utf-8 -*-
"""把 paper/manuscript.md 渲染进课程模板，生成 DOCX，并用 Word 导出 PDF。

用法：py src/build_docx.py [模板路径]
默认模板：../AI4S_Research_Template.docx（相对本仓库根目录）
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from docx_richtext import REF_HEAD, add_bookmark, add_rich_runs
from docx_equation import insert_equation

import params as P

DOCX_OUT = P.ROOT / "paper" / "AI4S论文-MLP多层介质薄膜光谱预测与辅助设计.docx"
PDF_OUT = DOCX_OUT.with_suffix(".pdf")
CENTER = WD_ALIGN_PARAGRAPH.CENTER
FIG_DPI = 300          # 图件保存 DPI，用于按原始尺寸插入


def set_two_columns(section, space_twips=280):
    """把某一节设为双栏（用于参考文献）。"""
    from docx.oxml.ns import qn
    sectPr = section._sectPr
    for old_cols in sectPr.findall(qn("w:cols")):
        sectPr.remove(old_cols)
    cols = sectPr.makeelement(qn("w:cols"), {qn("w:num"): "2", qn("w:space"): str(space_twips)})
    sectPr.append(cols)


def find_paragraph(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def cut_body_after(doc, anchor):
    body = doc.element.body
    el = anchor._element
    started = False
    for child in list(body):
        if child is el:
            started = True
            continue
        if not started or child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def set_text(par, text, bold=False):
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    run = par.add_run(text)
    run.bold = bold
    return par


def body_paragraph(doc, text, indent=True, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = 1.06          # 模板为 1.179，略微收紧以控制在 6 页内
    if indent:
        pf.first_line_indent = Pt(2 * size)
    pf.space_after = Pt(1.0)
    add_rich_runs(p, text, Pt(size))
    return p


def caption_paragraph(doc, text, size=8.5, keep=False):
    """图题/表题：居中，'**图 N**' 加粗。"""
    p = doc.add_paragraph()
    p.alignment = CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    if keep:
        p.paragraph_format.keep_with_next = True
    m = re.match(r"^(\*\*[^*]+\*\*)(.*)$", text)
    head, tail = (m.group(1).strip("*"), m.group(2)) if m else ("", text)
    r1 = p.add_run(head)
    r1.bold = True
    r1.font.size = Pt(size)
    add_rich_runs(p, tail, Pt(size))
    return p


def figure_block(doc, name):
    """按图片原始尺寸 1:1 插入，图内字号即设计字号；图题与本段绑定。"""
    from PIL import Image
    with Image.open(P.FIG_DIR / name) as im:
        width = im.width / FIG_DPI * 2.54      # 设计时的物理宽度（cm）
    p = doc.add_paragraph()
    p.alignment = CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(P.FIG_DIR / name), width=Cm(width))
    return p


def add_caption_into(par, text, size=8.5):
    """把图题追加到图片所在段落（图片下方另起一行）。"""
    par.paragraph_format.keep_with_next = False
    run = par.add_run()
    run.add_break()
    m = re.match(r"^(\*\*[^*]+\*\*)(.*)$", text)
    head, tail = (m.group(1).strip("*"), m.group(2)) if m else ("", text)
    r1 = par.add_run(head)
    r1.bold = True
    r1.font.size = Pt(size)
    add_rich_runs(par, tail, Pt(size))     # 图题同样支持下标与引用
    return par


def add_formula(doc, name):
    """插入原生 Word 公式（OMML）：全部由 docx_equation 直接构造，不调外部命令。"""
    insert_equation(doc, name)
    return None


def set_three_line_table(table):
    """按中文论文惯例把表格设为三线表：顶线、栏目线、底线，无竖线。"""
    from docx.oxml.ns import qn
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for edge, sz in (("top", 12), ("bottom", 12),
                     ("left", 0), ("right", 0), ("insideH", 0), ("insideV", 0)):
        el = borders.makeelement(qn(f"w:{edge}"), {})
        if sz:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "none")
            el.set(qn("w:sz"), "0")
        borders.append(el)
    tblPr.append(borders)
    # 栏目线：表头行下框线
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        tcB = tcPr.makeelement(qn("w:tcBorders"), {})
        bottom = tcB.makeelement(qn("w:bottom"), {})
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")
        tcB.append(bottom)
        tcPr.append(tcB)


def add_table_caption(doc, text, size=8.5):
    return caption_paragraph(doc, text, size=size, keep=True)


def add_top5_table(doc):
    ds = json.loads((P.RES_DIR / "design_screening.json").read_text(encoding="utf-8"))
    rows = ds["final_top5"]
    t = doc.add_table(rows=1 + len(rows), cols=7)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_three_line_table(t)
    head = ["排名", "d₁ / nm", "d₂ / nm", "d₃ / nm", "d₄ / nm",
            "MLP 预测 R(480 nm)", "TMM 复核 R(480 nm)"]
    for j, h in enumerate(head):
        cell = t.cell(0, j)
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(8)
        cell.paragraphs[0].alignment = CENTER
    for i, row in enumerate(rows, start=1):
        vals = [str(row["final_rank"])] + [f"{x:.2f}" for x in row["d_nm"]] + \
               [f"{row['R_mlp_target']:.4f}", f"{row['R_tmm_target']:.4f}"]
        for j, v in enumerate(vals):
            cell = t.cell(i, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(v)
            run.font.size = Pt(7.5)
            cell.paragraphs[0].alignment = CENTER
    return t


def add_hyperlink(paragraph, url, text, size=9.5):
    """把 URL 写成真正的 Word 超链接（Ctrl+点击可跳转）。"""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    for el in (color, underline, sz):
        rPr.append(el)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return paragraph


def render_markdown(doc, md_text):
    lines = md_text.splitlines()
    i = next(k for k, l in enumerate(lines) if l.startswith("## 摘要"))
    buf: list[str] = []
    in_refs = False
    pending_fig = None

    def join_lines(parts):
        """按 Markdown 软换行拼接，并在中英/数字边界补一个空格。"""
        out = ""
        latin = set("0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz%×°")
        for part in parts:
            part = part.strip()
            if not part:
                continue
            if out:
                a, b = out[-1], part[0]
                need = (a in latin and "一" <= b <= "鿿") or                        ("一" <= a <= "鿿" and b in latin)
                if need:
                    out += " "
            out += part
        return out

    def flush():
        if buf:
            body_paragraph(doc, join_lines(buf))
            buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
            i += 1
            continue
        if s.startswith("## 摘要") or s.startswith("# "):
            flush()
            i += 1
            continue
        if s.startswith("{FIG:") or s.startswith("{{FIG:"):
            flush()
            pending_fig = s.split(":")[1].rstrip("}")
            i += 1
            continue
        if s.startswith("{{EQ:"):
            flush()
            add_formula(doc, s[5:].rstrip("}"))
            i += 1
            continue
        if s.startswith("{{TABLE:top5}}"):
            flush()
            add_top5_table(doc)
            i += 1
            continue
        if s.startswith("**图") and pending_fig:
            flush()
            par = figure_block(doc, pending_fig)
            add_caption_into(par, s)
            pending_fig = None
            i += 1
            continue
        if s.startswith("**表"):
            flush()
            add_table_caption(doc, s)
            i += 1
            continue
        if s.startswith("### "):
            flush()
            h = doc.add_heading(s[4:], level=2)
            h.paragraph_format.space_before = Pt(5)
            h.paragraph_format.space_after = Pt(1)
            i += 1
            continue
        if s.startswith("## "):
            flush()
            title = s[3:]
            in_refs = title.startswith("参考文献")
            if in_refs:
                set_two_columns(doc.add_section(WD_SECTION.CONTINUOUS))
            h = doc.add_heading(title, level=1)
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            i += 1
            continue
        if s.startswith("**关键词**"):
            flush()
            body_paragraph(doc, s.replace("**", ""), indent=False)
            i += 1
            continue
        if in_refs and REF_HEAD.match(s):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(15)
            p.paragraph_format.first_line_indent = Pt(-15)
            p.paragraph_format.space_after = Pt(0.2)
            num = REF_HEAD.match(s).group(1)
            body = REF_HEAD.sub("", s)
            r1 = p.add_run("[")
            r2 = p.add_run(num)
            r3 = p.add_run("] ")
            r4 = p.add_run(body)
            for r in (r1, r2, r3, r4):
                r.font.size = Pt(7.0)
            start, end = add_bookmark(p, "_Ref" + num)   # 供正文交叉引用
            p._p.insert(list(p._p).index(r2._r), start)
            p._p.insert(list(p._p).index(r2._r) + 1, end)
            i += 1
            continue
        if line.startswith("    "):
            flush()
            body_paragraph(doc, s, indent=False, size=9, align=CENTER)
            i += 1
            continue
        if s.startswith("https://"):
            flush()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1.0)
            add_hyperlink(p, s, s, size=9.5)
            i += 1
            continue
        buf.append(s)
        i += 1
    flush()


def main():
    template = Path(sys.argv[1]) if len(sys.argv) > 1 else P.ROOT.parent / "AI4S_Research_Template.docx"
    if not template.exists():
        raise SystemExit(f"模板不存在：{template}")
    md_path = P.ROOT / "paper" / "manuscript.md"
    if not md_path.exists():
        raise SystemExit("请先运行 build_values.py 生成 manuscript.md")

    doc = Document(str(template))
    pl = find_paragraph(doc, "姓名：")
    set_text(pl, f"姓名：{P.STUDENT_NAME.strip() or '__________'}    学号：{P.STUDENT_ID}")
    pt = find_paragraph(doc, "λtarget")
    set_text(pt, f"λtarget：{int(P.LAMBDA_TARGET)} nm    seed：{P.SEED}    "
                 f"design_seed：{P.DESIGN_SEED}")
    anchor = find_paragraph(doc, "Abstract")
    set_text(anchor, "摘要")
    cut_body_after(doc, anchor)

    render_markdown(doc, md_path.read_text(encoding="utf-8"))
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX_OUT))
    print("DOCX:", DOCX_OUT)

    try:
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        d = word.Documents.Open(str(DOCX_OUT))
        d.SaveAs(str(PDF_OUT), FileFormat=17)
        pages = d.ComputeStatistics(2)
        d.Close(False)
        word.Quit()
        print("PDF:", PDF_OUT, "| 页数:", pages)
    except Exception as e:
        print("PDF 导出失败（DOCX 已生成）:", e)


if __name__ == "__main__":
    main()
